"""Excel report writer using the public openpyxl package."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from engine import PALETTE


def write_report(path, tables):
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in tables.items():
        ws = wb.create_sheet(name)
        ws.append(list(headers))
        for row in rows:
            ws.append([v.item() if hasattr(v, 'item') and callable(v.item) else v for v in row])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 29
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='17324D')
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(vertical='center')
        for j, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(36, max(15, len(str(header)) + 3))
        if name == 'TodayTrade' and rows:
            signal_col = headers.index('Signal') + 1
            for row_idx in range(2, ws.max_row + 1):
                signal = ws.cell(row_idx, signal_col).value
                if signal in PALETTE:
                    bg, fg = PALETTE[signal]
                    for col_idx in (1, signal_col):
                        cell = ws.cell(row_idx, col_idx)
                        cell.fill = PatternFill('solid', fgColor=bg.lstrip('#'))
                        cell.font = Font(bold=True, color=fg.lstrip('#'))
        if name == 'AnnualVsSPY' and rows:
            col = get_column_letter(headers.index('Excess pp') + 1)
            target = f'{col}2:{col}{ws.max_row}'
            ws.conditional_formatting.add(target, CellIsRule(operator='greaterThanOrEqual', formula=['0'], fill=PatternFill('solid', fgColor='C6EFCE')))
            ws.conditional_formatting.add(target, CellIsRule(operator='lessThan', formula=['0'], fill=PatternFill('solid', fgColor='F8CBAD')))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
