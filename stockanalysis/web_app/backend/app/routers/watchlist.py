"""Watchlist Router with Shariah-compliance validation and Excel import/export."""
from io import BytesIO
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.database import get_db
from app.models import WatchlistStock
from app.schemas import WatchlistStockCreate, WatchlistStockUpdate
from app.config import DEFAULT_UNIVERSE_PATH

router = APIRouter(prefix="/api/watchlist", tags=["Watchlist"])


@router.get("")
def get_watchlist(db: Session = Depends(get_db)):
    """Retrieve all stocks in the active universe."""
    stocks = db.query(WatchlistStock).order_by(WatchlistStock.ticker).all()
    return [{
        "ticker": s.ticker,
        "name": s.name,
        "enabled": s.enabled,
        "shariah_status": s.shariah_status,
        "shariah_source": s.shariah_source,
        "shariah_date": s.shariah_date,
        "sector": s.sector,
        "notes": s.notes,
        "updated_at": str(s.updated_at) if s.updated_at else None
    } for s in stocks]


@router.post("")
def add_watchlist_stock(stock_in: WatchlistStockCreate, db: Session = Depends(get_db)):
    """Add a stock to the watchlist."""
    ticker = stock_in.ticker.strip().upper()
    existing = db.query(WatchlistStock).filter(WatchlistStock.ticker == ticker).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Stock {ticker} already exists in watchlist")

    # Shariah compliance integrity check
    if stock_in.shariah_status == "VERIFIED":
        if not stock_in.shariah_source or not stock_in.shariah_date:
            raise HTTPException(
                status_code=400,
                detail="Cannot claim Shariah VERIFIED without a documented verification source and verification date."
            )

    stock = WatchlistStock(
        ticker=ticker,
        name=stock_in.name or ticker,
        enabled=stock_in.enabled,
        shariah_status=stock_in.shariah_status or "NOT VERIFIED",
        shariah_source=stock_in.shariah_source or "",
        shariah_date=stock_in.shariah_date or "",
        notes=stock_in.notes or ""
    )
    db.add(stock)
    db.commit()
    db.refresh(stock)
    return {"message": f"Added {ticker}", "stock": stock.ticker}


@router.put("/{ticker}")
def update_watchlist_stock(ticker: str, update_in: WatchlistStockUpdate, db: Session = Depends(get_db)):
    """Update stock configuration (enable/disable, Shariah status, notes)."""
    ticker = ticker.strip().upper()
    stock = db.query(WatchlistStock).filter(WatchlistStock.ticker == ticker).first()
    if not stock:
        raise HTTPException(status_code=404, detail=f"Stock {ticker} not found")

    if update_in.shariah_status is not None:
        target_status = update_in.shariah_status
        target_source = update_in.shariah_source if update_in.shariah_source is not None else stock.shariah_source
        target_date = update_in.shariah_date if update_in.shariah_date is not None else stock.shariah_date

        if target_status == "VERIFIED" and (not target_source or not target_date):
            raise HTTPException(
                status_code=400,
                detail="Cannot claim Shariah VERIFIED without a documented verification source and verification date."
            )
        stock.shariah_status = target_status

    if update_in.shariah_source is not None:
        stock.shariah_source = update_in.shariah_source
    if update_in.shariah_date is not None:
        stock.shariah_date = update_in.shariah_date
    if update_in.enabled is not None:
        stock.enabled = update_in.enabled
    if update_in.name is not None:
        stock.name = update_in.name
    if update_in.notes is not None:
        stock.notes = update_in.notes

    db.commit()
    return {"message": f"Updated {ticker}", "enabled": stock.enabled, "shariah_status": stock.shariah_status}


@router.delete("/{ticker}")
def delete_watchlist_stock(ticker: str, db: Session = Depends(get_db)):
    """Remove a stock from the watchlist."""
    ticker = ticker.strip().upper()
    stock = db.query(WatchlistStock).filter(WatchlistStock.ticker == ticker).first()
    if not stock:
        raise HTTPException(status_code=404, detail=f"Stock {ticker} not found")
    db.delete(stock)
    db.commit()
    return {"message": f"Removed {ticker}"}


@router.post("/import-default")
def import_default_universe(db: Session = Depends(get_db)):
    """Import initial universe from Stock_Universe.xlsx."""
    if not DEFAULT_UNIVERSE_PATH.exists():
        raise HTTPException(status_code=404, detail="Default Stock_Universe.xlsx not found")
    df = pd.read_excel(DEFAULT_UNIVERSE_PATH, sheet_name="Stocks")
    count = 0
    for _, r in df.iterrows():
        ticker = str(r["Ticker"]).strip().upper()
        if not ticker or ticker == "SPY":
            continue
        enabled_val = str(r.get("Enabled", "Yes")).strip().lower() in ("yes", "y", "1", "true")
        existing = db.query(WatchlistStock).filter(WatchlistStock.ticker == ticker).first()
        if not existing:
            stock = WatchlistStock(
                ticker=ticker,
                name=ticker,
                enabled=enabled_val,
                shariah_status="NOT VERIFIED",
                notes=str(r.get("Notes", "")) if pd.notna(r.get("Notes")) else ""
            )
            db.add(stock)
            count += 1
    db.commit()
    return {"message": f"Imported {count} stocks from default universe"}


@router.get("/export-excel")
def export_universe_excel(db: Session = Depends(get_db)):
    """Export the stock universe to an Excel file formatted like Stock_Universe.xlsx."""
    stocks = db.query(WatchlistStock).order_by(WatchlistStock.ticker).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Stocks"
    headers = ["Ticker", "Enabled", "Shariah Status", "Verification Source", "Verification Date", "Notes"]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(bold=True, color="FFFFFF")

    # Add SPY benchmark row first
    ws.append(["SPY", "Yes", "Benchmark only", "", "", "SPY benchmark proxy"])

    for s in stocks:
        ws.append([
            s.ticker,
            "Yes" if s.enabled else "No",
            s.shariah_status,
            s.shariah_source,
            s.shariah_date,
            s.notes
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Stock_Universe_Export.xlsx"}
    )
